import streamlit as st
import google.generativeai as genai
import io

# Optional file-parsing libraries (loaded lazily, only when needed)
def extract_text_from_pdf(file_bytes):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text

def extract_text_from_docx(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs)

def extract_text_from_upload(uploaded_file):
    """Returns plain text content from an uploaded PDF, DOCX, or TXT file."""
    if uploaded_file is None:
        return ""
    name = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    try:
        if name.endswith(".pdf"):
            return extract_text_from_pdf(file_bytes)
        elif name.endswith(".docx"):
            return extract_text_from_docx(file_bytes)
        elif name.endswith(".txt"):
            return file_bytes.decode("utf-8", errors="ignore")
        else:
            return ""
    except Exception as e:
        st.warning(f"Dili nabasa ang file ({uploaded_file.name}): {e}")
        return ""


# ----------------------------------------------------------------------------
# Page setup
# ----------------------------------------------------------------------------
st.set_page_config(page_title="AI Daily Lesson Log (DLL) Planner", layout="wide")

st.markdown(
    """
    <div style="background-color:#1A56DB;padding:18px 24px;border-radius:8px;margin-bottom:20px;">
        <h2 style="color:white;margin:0;">📘 AI Daily Lesson Log (DLL) Planner</h2>
        <p style="color:#D6E4FF;margin:0;">DepEd MATATAG-aligned • AI-generated lesson plans</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------------
# Sidebar: API Key
# ----------------------------------------------------------------------------
with st.sidebar:
    st.header("⚙️ Setup")
    api_key = st.text_input(
        "Gemini API Key",
        type="password",
        help="Kuhaa sa https://aistudio.google.com/apikey",
        value=st.secrets.get("GEMINI_API_KEY", "") if hasattr(st, "secrets") else "",
    )
    st.caption("Ang imong API key dili na-save bisan asa. Gamiton ra sa kasamtangang session.")
    st.divider()
    st.caption("Developed with reference to the AI DLL Planner Gem.")

if api_key:
    genai.configure(api_key=api_key)

MODEL_NAME = "gemini-3.5-flash"  # current GA fast model as of mid-2026

# ----------------------------------------------------------------------------
# Main layout: two columns
# ----------------------------------------------------------------------------
col1, col2 = st.columns([1, 1.3])

with col1:
    st.subheader("📝 Lesson Details")

    language = st.radio("Language / Wika *", ["English", "Filipino"], horizontal=True)

    input_mode = st.radio(
        "Input Mode *",
        ["A. Manual Input", "B. Upload Exemplar", "C. Upload Other Sources"],
        horizontal=True,
    )

    exemplar_text = ""
    other_sources_text = ""

    if input_mode == "B. Upload Exemplar":
        exemplar_file = st.file_uploader(
            "I-upload ang sample/exemplar DLL (PDF, DOCX, or TXT)",
            type=["pdf", "docx", "txt"],
            key="exemplar",
        )
        if exemplar_file:
            exemplar_text = extract_text_from_upload(exemplar_file)
            if exemplar_text:
                st.success(f"Na-load ang exemplar: {exemplar_file.name}")

    elif input_mode == "C. Upload Other Sources":
        other_files = st.file_uploader(
            "I-upload ang reference materials (curriculum guide, textbook excerpt, atbp.)",
            type=["pdf", "docx", "txt"],
            accept_multiple_files=True,
            key="other_sources",
        )
        if other_files:
            texts = []
            for f in other_files:
                t = extract_text_from_upload(f)
                if t:
                    texts.append(f"--- {f.name} ---\n{t}")
            other_sources_text = "\n\n".join(texts)
            if other_sources_text:
                st.success(f"Na-load ang {len(other_files)} file(s).")

    learning_area = st.text_input("Learning Area/s *", placeholder="e.g. Science 7")

    c1, c2 = st.columns(2)
    with c1:
        teacher_name = st.text_input("Designed by Teacher/s *", placeholder="e.g. Juan Dela Cruz")
    with c2:
        teacher_role = st.text_input("Teacher Role", placeholder="e.g. Teacher III")

    c3, c4 = st.columns(2)
    with c3:
        checker_name = st.text_input("Checked by *", placeholder="e.g. Maria Santos")
    with c4:
        checker_role = st.text_input("Checker Role", placeholder="e.g. Master Teacher II")

    grade_section = st.text_input("Grade Level and Section *", placeholder="e.g. Grade 7 - Archimedes")

    num_days = st.selectbox(
        "No. of DLLs to Generate *",
        options=[1, 2, 3, 4, 5],
        index=3,
        format_func=lambda x: f"{x} Day{'s' if x > 1 else ''} ({x} Separate DLL{'s' if x > 1 else ''})",
    )
    st.caption("ℹ️ Maghimo ang AI og separate table alang sa kada adlaw.")

    competency = st.text_area(
        "Learning Competency/ies *",
        placeholder="e.g. Identify parts of the microscope and their functions. (S7LT-IIa-1)",
        height=80,
    )

    objectives = st.text_area(
        "Learning Objectives (Optional)",
        placeholder="1. Identify the major parts of a compound light microscope.\n2. Classify the parts according to their function.",
        height=100,
    )

    generate_btn = st.button("✏️ Generate Lesson Plan", type="primary", use_container_width=True)

with col2:
    st.subheader("📄 Generated Lesson Plan(s)")

    output_area = st.container()

    if generate_btn:
        # --- Validation ---
        missing = []
        if not api_key:
            missing.append("Gemini API Key (sa sidebar)")
        if not learning_area:
            missing.append("Learning Area/s")
        if not teacher_name:
            missing.append("Designed by Teacher/s")
        if not checker_name:
            missing.append("Checked by")
        if not grade_section:
            missing.append("Grade Level and Section")
        if not competency:
            missing.append("Learning Competency/ies")
        if input_mode == "B. Upload Exemplar" and not exemplar_text:
            missing.append("Upload Exemplar file (o pilia ang lain nga Input Mode)")

        if missing:
            output_area.error(
                "⚠️ Palihog kompletoha ning mga kinahanglanon nga field:\n\n"
                + "\n".join(f"- {m}" for m in missing)
            )
        else:
            with output_area:
                with st.spinner("Binubuo ang imong lesson plan..."):
                    try:
                        model = genai.GenerativeModel(MODEL_NAME)

                        context_blocks = []
                        if input_mode == "B. Upload Exemplar" and exemplar_text:
                            context_blocks.append(
                                "Gamita kini nga exemplar/sample DLL isip GABAY sa FORMAT, ESTILO, "
                                "ug LEBEL SA DETALYE (dili kopyahon ang sulod, format lang sundon):\n"
                                + exemplar_text[:6000]
                            )
                        if input_mode == "C. Upload Other Sources" and other_sources_text:
                            context_blocks.append(
                                "Gamita kini nga reference materials isip BASEHANAN sa sulod sa leksyon "
                                "(curriculum guide, textbook excerpt, ubp.):\n"
                                + other_sources_text[:6000]
                            )
                        extra_context = "\n\n".join(context_blocks)

                        prompt = f"""
You are an expert Filipino public school teacher creating DepEd MATATAG-aligned Daily Lesson Logs (DLL).

Generate a {num_days}-day Daily Lesson Log (DLL) with the following details:

- Language to write the lesson plan in: {language}
- Learning Area/s: {learning_area}
- Designed by Teacher/s: {teacher_name} ({teacher_role or "Teacher"})
- Checked by: {checker_name} ({checker_role or "Master Teacher"})
- Grade Level and Section: {grade_section}
- Learning Competency/ies: {competency}
- Learning Objectives: {objectives if objectives else "Derive appropriate objectives from the competency above."}

{extra_context}

INSTRUCTIONS:
1. Create {num_days} SEPARATE tables, one per day (Day 1 to Day {num_days}).
2. Each table must use the standard DepEd DLL rows in this exact order:
   - Balik-aral (Review)
   - Pagganyak (Motivation)
   - Paglalahad/Talakayan (Presentation/Discussion)
   - Paglalapat (Application)
   - Paglalahat (Generalization)
   - Pagtataya (Assessment/Evaluation)
3. Each cell should contain specific, classroom-ready content (actual activities, sample questions, or instructions) — not generic placeholders.
4. Make sure content builds progressively across days when relevant (e.g. Day 1 review feeds into Day 2).
5. Format the entire output in clean Markdown, with a level-2 heading for each day (## Day 1, ## Day 2, etc.) followed by a Markdown table with two columns: "Yugto ng Aralin" and the content for that day.
6. At the top, include a header block with: Learning Area, Teacher, Checked by, Grade & Section, and Learning Competency/ies.
7. Do not include any commentary outside the lesson plan itself.
"""
                        response = model.generate_content(prompt)
                        st.markdown(response.text)

                        st.download_button(
                            "⬇️ Download as Markdown",
                            data=response.text,
                            file_name=f"DLL_{learning_area.replace(' ', '_')}.md",
                            mime="text/markdown",
                        )

                    except Exception as e:
                        st.error(f"Naay sayop nga nahitabo: {e}")
    else:
        output_area.info(
            "🤖 Ilagay ang detalye sa wala para makabuo og lesson plan.\n\n"
            "Maghimo ang AI og separate, complete nga DLL kada adlaw, base sa DepEd guidelines."
        )

st.divider()
st.caption("⚠️ AI-generated content — palihog susiha ug i-adjust base sa imong klase ug konteksto sa ⁠eskwelahan.")
